import os
from datetime import datetime, date, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO

import pymysql
from pymysql.cursors import DictCursor
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from flask_wtf.csrf import CSRFProtect
from dotenv import load_dotenv
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

load_dotenv()
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-only-change-this-secret")
app.config["WTF_CSRF_TIME_LIMIT"] = None
csrf = CSRFProtect(app)

def db():
    return pymysql.connect(
        host=os.getenv("DB_HOST", "localhost"),
        user=os.getenv("DB_USER", "root"),
        password=os.getenv("DB_PASSWORD", ""),
        database=os.getenv("DB_NAME", "moneymate"),
        cursorclass=DictCursor,
        autocommit=False,
        charset="utf8mb4",
    )

def query_all(sql, args=()):
    conn = db()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, args)
            return cur.fetchall()
    finally:
        conn.close()

def query_one(sql, args=()):
    rows = query_all(sql, args)
    return rows[0] if rows else None

def execute(sql, args=()):
    conn = db()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, args)
            result = cur.lastrowid
        conn.commit()
        return result
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

def current_user_id():
    return int(session["user_id"])

def money(value):
    try:
        amount = Decimal(str(value))
        if not amount.is_finite() or amount <= 0:
            raise ValueError
        return amount.quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Enter an amount greater than zero.")

def parse_datetime(value, time_value):
    try:
        return datetime.strptime(f"{value} {time_value}", "%Y-%m-%d %H:%M").strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        raise ValueError("Choose a valid date and time.")

def user_accounts():
    return query_all("SELECT * FROM accounts WHERE user_id=%s AND archived_at IS NULL ORDER BY account_name", (current_user_id(),))

def get_balance(account_id, conn=None):
    own_conn = conn is None
    conn = conn or db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT opening_balance FROM accounts WHERE id=%s AND user_id=%s", (account_id, current_user_id()))
            account = cur.fetchone()
            if not account:
                raise ValueError("Account not found.")
            opening = Decimal(str(account["opening_balance"]))
            cur.execute("""SELECT
                COALESCE(SUM(CASE WHEN transaction_type='income' AND account_id=%s THEN amount
                                  WHEN transaction_type='transfer' AND to_account_id=%s THEN amount
                                  WHEN transaction_type='expense' AND account_id=%s THEN -amount
                                  WHEN transaction_type='transfer' AND from_account_id=%s THEN -amount
                                  ELSE 0 END),0) AS delta
                FROM transactions WHERE user_id=%s""",
                (account_id, account_id, account_id, account_id, current_user_id()))
            delta = Decimal(str(cur.fetchone()["delta"]))
            return opening + delta
    finally:
        if own_conn:
            conn.close()

def period_bounds(period, selected_month=None):
    now = datetime.now()
    if period == "week":
        start = (now - timedelta(days=now.weekday())).date()
        end = start + timedelta(days=7)
    elif period == "year":
        start = date(now.year, 1, 1)
        end = date(now.year + 1, 1, 1)
    elif period == "month" and selected_month:
        start = datetime.strptime(selected_month, "%Y-%m").date().replace(day=1)
        end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    else:
        start = now.date()
        end = start + timedelta(days=1)
    return start, end

@app.context_processor
def inject_globals():
    return {"app_name": "MoneyMate", "current_user": session.get("phone_number")}

@app.route("/")
def index():
    return redirect(url_for("dashboard") if session.get("user_id") else url_for("login"))

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        phone = request.form.get("phone_number", "").strip()
        password = request.form.get("password", "")
        if len(phone) < 8 or len(phone) > 20 or len(password) < 8:
            flash("Enter a valid phone number and a password of at least 8 characters.", "error")
            return render_template("register.html")
        try:
            uid = execute("INSERT INTO users(phone_number,password_hash) VALUES(%s,%s)", (phone, generate_password_hash(password)))
            execute("INSERT INTO accounts(user_id,account_name,account_type,opening_balance) VALUES(%s,'Cash','Cash',0)", (uid,))
            session["user_id"], session["phone_number"] = uid, phone
            flash("Account created. You can start tracking your money.", "success")
            return redirect(url_for("dashboard"))
        except pymysql.err.IntegrityError:
            flash("That phone number is already registered.", "error")
    return render_template("register.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        phone = request.form.get("phone_number", "").strip()
        password = request.form.get("password", "")
        user = query_one("SELECT * FROM users WHERE phone_number=%s", (phone,))
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"], session["phone_number"] = user["id"], user["phone_number"]
            return redirect(url_for("dashboard"))
        flash("Invalid phone number or password.", "error")
    return render_template("login.html")

@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/dashboard")
@login_required
def dashboard():
    today = date.today()
    month_start = today.replace(day=1)
    summary = query_one("""SELECT
        COALESCE(SUM(CASE WHEN transaction_type='income' AND DATE(transaction_datetime)=%s THEN amount ELSE 0 END),0) AS daily_income,
        COALESCE(SUM(CASE WHEN transaction_type='expense' AND DATE(transaction_datetime)=%s THEN amount ELSE 0 END),0) AS daily_expense,
        COALESCE(SUM(CASE WHEN transaction_type='income' AND transaction_datetime >= %s THEN amount ELSE 0 END),0) AS monthly_income,
        COALESCE(SUM(CASE WHEN transaction_type='expense' AND transaction_datetime >= %s THEN amount ELSE 0 END),0) AS monthly_expense,
        COALESCE(SUM(CASE WHEN transaction_type='transfer' AND transaction_datetime >= %s THEN amount ELSE 0 END),0) AS monthly_transfers
        FROM transactions WHERE user_id=%s""", (today, today, month_start, month_start, month_start, current_user_id()))
    recent = query_all("SELECT * FROM transactions WHERE user_id=%s ORDER BY transaction_datetime DESC,id DESC LIMIT 8", (current_user_id(),))
    accounts = user_accounts()
    total_balance = sum((get_balance(a["id"]) for a in accounts), Decimal("0.00"))
    category_rows = query_all("""SELECT category, SUM(amount) AS total FROM transactions
        WHERE user_id=%s AND transaction_type='expense' AND transaction_datetime >= %s
        GROUP BY category ORDER BY total DESC LIMIT 6""", (current_user_id(), month_start))
    return render_template("dashboard.html", summary=summary, recent=recent, accounts=accounts,
                           total_balance=total_balance, category_rows=category_rows, today=today)

@app.route("/transactions")
@login_required
def transactions():
    kind = request.args.get("type", "")
    month = request.args.get("month", date.today().strftime("%Y-%m"))
    category = request.args.get("category", "")
    account = request.args.get("account", "")
    search = request.args.get("q", "").strip()
    start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    sql = "SELECT * FROM transactions WHERE user_id=%s AND transaction_datetime >= %s AND transaction_datetime < %s"
    args = [current_user_id(), start, end]
    if kind in ("income", "expense", "transfer"):
        sql += " AND transaction_type=%s"; args.append(kind)
    if category:
        sql += " AND category=%s"; args.append(category)
    if account:
        sql += " AND (account_id=%s OR from_account_id=%s OR to_account_id=%s)"; args.extend([account, account, account])
    if search:
        sql += " AND (COALESCE(description,'') LIKE %s OR COALESCE(note,'') LIKE %s)"; args.extend([f"%{search}%", f"%{search}%"])
    sql += " ORDER BY transaction_datetime DESC,id DESC"
    rows = query_all(sql, tuple(args))
    totals = query_one("""SELECT
        COALESCE(SUM(CASE WHEN transaction_type='income' THEN amount ELSE 0 END),0) AS income,
        COALESCE(SUM(CASE WHEN transaction_type='expense' THEN amount ELSE 0 END),0) AS expense,
        COALESCE(SUM(CASE WHEN transaction_type='transfer' THEN amount ELSE 0 END),0) AS transfers
        FROM transactions WHERE user_id=%s AND transaction_datetime >= %s AND transaction_datetime < %s""", (current_user_id(), start, end))
    return render_template("transactions.html", rows=rows, month=month, kind=kind, category=category,
                           account=account, search=search, accounts=user_accounts(), totals=totals)

def save_transaction(kind, tx_id=None):
    amount = money(request.form.get("amount"))
    tx_date = parse_datetime(request.form.get("date"), request.form.get("time"))
    category = request.form.get("category", "").strip()
    account_id = request.form.get("account_id") or None
    description = request.form.get("description", "").strip() or None
    note = request.form.get("note", "").strip() or None
    allowed = {
        "income": {"Allowance", "Salary", "Petty Cash", "Bonus", "Other"},
        "expense": {"Food", "Social Life", "Pets", "Transport", "Culture", "Household", "Apparel", "Beauty", "Health", "Education", "Gift", "Other"},
    }
    if category not in allowed[kind]:
        raise ValueError("Choose a valid category.")
    if not account_id:
        raise ValueError("Choose an account.")
    account = query_one("SELECT id FROM accounts WHERE id=%s AND user_id=%s AND archived_at IS NULL", (account_id, current_user_id()))
    if not account:
        raise ValueError("Choose one of your active accounts.")
    if kind == "expense" and get_balance(account_id) < amount:
        raise ValueError("This account does not have enough balance for that expense.")
    if tx_id:
        existing = query_one("SELECT id FROM transactions WHERE id=%s AND user_id=%s AND transaction_type=%s", (tx_id, current_user_id(), kind))
        if not existing:
            raise ValueError("Transaction not found.")
        execute("""UPDATE transactions SET amount=%s,category=%s,account_id=%s,description=%s,note=%s,transaction_datetime=%s
                   WHERE id=%s AND user_id=%s""", (amount, category, account_id, description, note, tx_date, tx_id, current_user_id()))
    else:
        execute("""INSERT INTO transactions(user_id,transaction_type,amount,category,account_id,description,note,transaction_datetime)
                   VALUES(%s,%s,%s,%s,%s,%s,%s,%s)""",
                (current_user_id(), kind, amount, category, account_id, description, note, tx_date))

@app.route("/add/<kind>", methods=["GET", "POST"])
@login_required
def add_transaction(kind):
    if kind not in ("income", "expense"):
        return redirect(url_for("dashboard"))
    tx_id = request.args.get("edit", type=int)
    existing = query_one("SELECT * FROM transactions WHERE id=%s AND user_id=%s AND transaction_type=%s", (tx_id, current_user_id(), kind)) if tx_id else None
    if request.method == "POST":
        try:
            save_transaction(kind, tx_id)
            flash(f"{kind.title()} {'updated' if tx_id else 'saved'} successfully.", "success")
            return redirect(url_for("transactions"))
        except (ValueError, pymysql.MySQLError) as exc:
            flash(str(exc) if isinstance(exc, ValueError) else "Could not save this transaction. Please check your database settings.", "error")
    now = datetime.now()
    return render_template("transaction_form.html", kind=kind, existing=existing, accounts=user_accounts(),
                           now=now, categories=(["Allowance","Salary","Petty Cash","Bonus","Other"] if kind=="income" else ["Food","Social Life","Pets","Transport","Culture","Household","Apparel","Beauty","Health","Education","Gift","Other"]))

@app.route("/transfer", methods=["GET", "POST"])
@login_required
def transfer():
    now = datetime.now()
    if request.method == "POST":
        try:
            amount = money(request.form.get("amount"))
            tx_date = parse_datetime(request.form.get("date"), request.form.get("time"))
            from_id, to_id = request.form.get("from_account_id"), request.form.get("to_account_id")
            note = request.form.get("note", "").strip() or None
            if not from_id or not to_id or from_id == to_id:
                raise ValueError("Choose two different accounts.")
            conn = db()
            try:
                with conn.cursor() as cur:
                    cur.execute("SELECT id FROM accounts WHERE id=%s AND user_id=%s AND archived_at IS NULL FOR UPDATE", (from_id, current_user_id()))
                    src = cur.fetchone()
                    cur.execute("SELECT id FROM accounts WHERE id=%s AND user_id=%s AND archived_at IS NULL FOR UPDATE", (to_id, current_user_id()))
                    dest = cur.fetchone()
                    if not src or not dest:
                        raise ValueError("Choose valid accounts.")
                    if get_balance(from_id, conn) < amount:
                        raise ValueError("The source account does not have enough balance.")
                    cur.execute("""INSERT INTO transactions(user_id,transaction_type,amount,category,from_account_id,to_account_id,note,transaction_datetime)
                                    VALUES(%s,'transfer',%s,'Transfer',%s,%s,%s,%s)""",
                                (current_user_id(), amount, from_id, to_id, note, tx_date))
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
            flash("Transfer saved successfully.", "success")
            return redirect(url_for("transactions", month=tx_date[:7]))
        except (ValueError, pymysql.MySQLError) as exc:
            flash(str(exc) if isinstance(exc, ValueError) else "Could not save the transfer.", "error")
    return render_template("transfer.html", accounts=user_accounts(), now=now)

@app.route("/delete/<int:tx_id>", methods=["POST"])
@login_required
def delete_transaction(tx_id):
    row = query_one("SELECT * FROM transactions WHERE id=%s AND user_id=%s", (tx_id, current_user_id()))
    if not row:
        flash("Transaction not found.", "error")
        return redirect(url_for("transactions"))
    if row["transaction_type"] == "transfer":
        flash("Transfer deletion is disabled in this starter build to protect account balances. Add a transfer reversal workflow before enabling it.", "error")
    else:
        execute("DELETE FROM transactions WHERE id=%s AND user_id=%s", (tx_id, current_user_id()))
        flash("Transaction deleted.", "success")
    return redirect(url_for("transactions"))

@app.route("/accounts", methods=["GET", "POST"])
@login_required
def accounts():
    if request.method == "POST":
        name = request.form.get("account_name", "").strip()
        kind = request.form.get("account_type", "Cash")
        try:
            opening = Decimal(request.form.get("opening_balance", "0"))
            if not name or len(name) > 80 or opening < 0 or kind not in ("Cash", "Bank Account", "Card"):
                raise ValueError
            execute("INSERT INTO accounts(user_id,account_name,account_type,opening_balance) VALUES(%s,%s,%s,%s)",
                    (current_user_id(), name, kind, opening))
            flash("Account added.", "success")
            return redirect(url_for("accounts"))
        except (ValueError, InvalidOperation, pymysql.MySQLError):
            flash("Enter a valid account name and opening balance.", "error")
    rows = user_accounts()
    for row in rows:
        row["balance"] = get_balance(row["id"])
    return render_template("accounts.html", accounts=rows)

@app.route("/notes", methods=["GET", "POST"])
@login_required
def notes():
    if request.method == "POST":
        title = request.form.get("title", "").strip()
        content = request.form.get("content", "").strip()
        if not title or not content:
            flash("Add both a title and note content.", "error")
        else:
            execute("INSERT INTO notes(user_id,title,content) VALUES(%s,%s,%s)", (current_user_id(), title[:120], content))
            flash("Note saved.", "success")
            return redirect(url_for("notes"))
    rows = query_all("SELECT * FROM notes WHERE user_id=%s ORDER BY updated_at DESC,id DESC", (current_user_id(),))
    return render_template("notes.html", notes=rows)

@app.route("/notes/delete/<int:note_id>", methods=["POST"])
@login_required
def delete_note(note_id):
    execute("DELETE FROM notes WHERE id=%s AND user_id=%s", (note_id, current_user_id()))
    flash("Note deleted.", "success")
    return redirect(url_for("notes"))

@app.route("/statistics")
@login_required
def statistics():
    period = request.args.get("period", "month")
    if period not in ("week", "month", "year"):
        period = "month"
    start, end = period_bounds(period, request.args.get("month"))
    rows = query_all("""SELECT transaction_type,category,SUM(amount) AS total FROM transactions
        WHERE user_id=%s AND transaction_datetime >= %s AND transaction_datetime < %s
        GROUP BY transaction_type,category ORDER BY transaction_type, total DESC""", (current_user_id(), start, end))
    daily = query_all("""SELECT DATE(transaction_datetime) AS day,
        SUM(CASE WHEN transaction_type='income' THEN amount ELSE 0 END) AS income,
        SUM(CASE WHEN transaction_type='expense' THEN amount ELSE 0 END) AS expense,
        SUM(CASE WHEN transaction_type='transfer' THEN amount ELSE 0 END) AS transfers
        FROM transactions WHERE user_id=%s AND transaction_datetime >= %s AND transaction_datetime < %s
        GROUP BY DATE(transaction_datetime) ORDER BY day""", (current_user_id(), start, end))
    return render_template("statistics.html", period=period, rows=rows, daily=daily, start=start, end=end)

@app.route("/export")
@login_required
def export():
    month = request.args.get("month", date.today().strftime("%Y-%m"))
    start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    rows = query_all("""SELECT DATE_FORMAT(transaction_datetime,'%%Y-%%m-%%d') AS date,
        DATE_FORMAT(transaction_datetime,'%%H:%%i') AS time, transaction_type AS type, category, description,
        account_id, from_account_id, to_account_id, amount, note
        FROM transactions WHERE user_id=%s AND transaction_datetime >= %s AND transaction_datetime < %s
        ORDER BY transaction_datetime""", (current_user_id(), start, end))
    accounts_data = user_accounts()
    for r in rows:
        r["account"] = next((a["account_name"] for a in accounts_data if a["id"] == r["account_id"]), "")
        r["from_account"] = next((a["account_name"] for a in accounts_data if a["id"] == r["from_account_id"]), "")
        r["to_account"] = next((a["account_name"] for a in accounts_data if a["id"] == r["to_account_id"]), "")
        for k in ("account_id", "from_account_id", "to_account_id"):
            r.pop(k, None)
    wb = Workbook()
    ws = wb.active
    ws.title = "All Transactions"
    headers = ["Date","Time","Type","Category","Description","Account","From Account","To Account","Amount","Note"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h.lower().replace(" ", "_")) for h in headers])
    for sheet in [ws]:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25324B")
            cell.alignment = Alignment(vertical="center")
        for col in sheet.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 36)
    for title, tx_type in (("Income Summary","income"),("Expense Summary","expense"),("Transfers","transfer")):
        sh = wb.create_sheet(title)
        sh.append(["Category","Total"])
        totals = {}
        for r in rows:
            if r["type"] == tx_type:
                key = r["category"] or tx_type.title()
                totals[key] = totals.get(key, Decimal("0")) + Decimal(str(r["amount"]))
        for key, val in sorted(totals.items()):
            sh.append([key, float(val)])
        sh.freeze_panes = "A2"
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="25324B")
        sh.column_dimensions["A"].width = 24
        sh.column_dimensions["B"].width = 18
    sh = wb.create_sheet("Account Summary")
    sh.append(["Account","Type","Opening Balance","Current Balance"])
    for a in accounts_data:
        sh.append([a["account_name"], a["account_type"], float(a["opening_balance"]), float(get_balance(a["id"]))])
    for c in sh[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="25324B")
    for col in sh.columns:
        sh.column_dimensions[get_column_letter(col[0].column)].width = 22
    output = BytesIO()
    wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"MoneyMate_Transactions_{month}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.errorhandler(404)
def not_found(_):
    return render_template("error.html", code=404, message="The page you requested could not be found."), 404

@app.errorhandler(500)
def server_error(_):
    return render_template("error.html", code=500, message="Something went wrong. Check the server terminal for details."), 500

if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "0") == "1", port=int(os.getenv("PORT", "5000")))
